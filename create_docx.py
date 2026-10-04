import docx
import re
import os

def create_docx(input_file, output_file):
    doc = docx.Document()
    
    with open(input_file, 'r', encoding='utf-8') as f:
        content = f.read()

    # We want to extract lessons from '## الدرس الأول' to the end of lessons (before '## معايير التقييم لكل درس')
    start_idx = content.find('## الدرس الأول')
    end_idx = content.find('## معايير التقييم لكل درس')
    
    if start_idx != -1 and end_idx != -1:
        lessons_content = content[start_idx:end_idx]
    else:
        lessons_content = content

    # Add title
    doc.add_heading('محتوى دروس المنصة - الفصل الأول', 0)
    
    for line in lessons_content.split('\n'):
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(line[2:].strip(), style='List Bullet')
        elif re.match(r'^\d+\.\s', line):
            doc.add_paragraph(line[line.find(' ')+1:].strip(), style='List Number')
        elif line.strip() == '':
            continue
        elif line.startswith('**') and line.endswith('**'):
            p = doc.add_paragraph()
            p.add_run(line.strip('* ')).bold = True
        else:
            # Check for bold parts within the line
            parts = re.split(r'(\*\*.*?\*\*)', line)
            p = doc.add_paragraph()
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)

    # Set right-to-left for all paragraphs
    for p in doc.paragraphs:
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
        p.style.font.rtl = True
        
    doc.save(output_file)
    print(f'Saved to {output_file}')

if __name__ == "__main__":
    create_docx('lessons_documentation.md', 'Semester1_Lessons.docx')
